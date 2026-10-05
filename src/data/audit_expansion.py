"""Milestone 1C: inventory acquired data without modelling or merging label sources."""

import argparse
from collections import Counter
import json
from pathlib import Path
import struct
import xml.etree.ElementTree as ET
import zipfile

import numpy as np
import openpyxl
import pandas as pd

from .acquire_expansion import RAW, ROOT, digest

OUTPUT = ROOT / "data/processed/expansion_audit"
EVENTS = {
    "Bingol earthquake engineering report": "bingol_2003",
    "Damage caused by the 1999 Duzce Earthquake in Turkey": "duzce_1999",
    "Damage caused by the 1992 Erzincan Earthquake": "erzincan_1992",
    "Measures of the seismic vulnerability of reinforced concrete buildings in Haiti": "haiti_2010",
    "Learning from the 2007 Pisco Peru earthquake": "pisco_2007",
    "Seismic vulnerability of reinforced concrete structures affected by the 2008 Wenchuan earthquake": "wenchuan_2008",
}
RC_DAMAGE = {"N": 0, "None": 0, "L": 1, "Light": 1, "M": 2, "Moderate": 2,
             "S": 3, "Severe": 3, "C": 4}
ACI_DAMAGE = {"N": 0, "L": 1, "M": 2, "S": 3, "C": 4}


def serializable(value):
    if isinstance(value, dict):
        return {str(k): serializable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [serializable(v) for v in value]
    if value is pd.NA or (isinstance(value, (float, np.floating)) and not np.isfinite(value)):
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(serializable(value), indent=2, ensure_ascii=False, allow_nan=False) + "\n")


def counts(series):
    return series.astype("string").fillna("<missing>").replace("", "<missing>").value_counts().sort_index().to_dict()


def table_profile(frame):
    profile = {}
    for name in frame:
        values = frame[name]
        missing = values.isna() | values.astype("string").str.strip().eq("").fillna(False)
        profile[name] = {"dtype": str(values.dtype), "missing": int(missing.sum()),
                         "distinct_nonmissing": int(values.loc[~missing].nunique())}
    return {"rows": len(frame), "variables": len(frame.columns), "columns": profile}


def coordinate_profile(frame, lat="latitude", lon="longitude"):
    x = pd.to_numeric(frame[lat], errors="coerce")
    y = pd.to_numeric(frame[lon], errors="coerce")
    present = x.notna() & y.notna()
    valid = present & x.between(-90, 90) & y.between(-180, 180) & ~((x == 0) & (y == 0))
    pairs = pd.DataFrame({"latitude": x[valid], "longitude": y[valid]})
    return {"present_pairs": int(present.sum()), "missing_or_unparseable_pairs": int((~present).sum()),
            "invalid_present_pairs": int((present & ~valid).sum()), "valid_pairs": int(valid.sum()),
            "unique_valid_pairs": len(pairs.drop_duplicates()),
            "rows_in_duplicate_coordinate_groups": int(pairs.duplicated(keep=False).sum()),
            "bounds_lon_lat": [y[valid].min(), x[valid].min(), y[valid].max(), x[valid].max()]}


def parse_nso_ddi(path):
    root = ET.parse(path).getroot()
    namespace = {"d": root.tag.split("}")[0].lstrip("{")}

    def text(node, query):
        return node.findtext(query, default="", namespaces=namespace).strip()

    files = root.findall("d:fileDscr", namespace)
    building = next(f for f in files if text(f, "d:fileTxt/d:fileName").startswith("Building."))
    variables = []
    for var in root.findall(".//d:var", namespace):
        if building.get("ID") not in var.get("files", "").split():
            continue
        categories = [{"value": text(c, "d:catValu"), "label": text(c, "d:labl"),
                       "frequency": int(text(c, "d:catStat")) if text(c, "d:catStat") else None,
                       "missing": c.get("missing") == "Y"} for c in var.findall("d:catgry", namespace)]
        variables.append({"name": var.get("name"), "label": text(var, "d:labl"), "categories": categories,
                          "summary_statistics": {s.get("type"): (s.text or "").strip() for s in var.findall("d:sumStat", namespace)}})
    return {"advertised_rows": int(text(building, "d:fileTxt/d:dimensns/d:caseQnty")),
            "advertised_variables": int(text(building, "d:fileTxt/d:dimensns/d:varQnty")),
            "variables": variables, "evidence_kind": "Published DDI metadata, not locally inspected microdata"}


def audit_nepal(output):
    official = parse_nso_ddi(RAW / "audit_metadata/nso_69.xml")
    definitions = {v["name"]: v for v in official["variables"]}
    write_json(output / "nso_building_dictionary.json", official)
    pd.DataFrame([{"name": v["name"], "description": v["label"]} for v in official["variables"]]).to_csv(output / "nso_building_variables.csv", index=False)
    source = RAW / "nepal_2015/Impact_Buildings_Detailed.csv"
    geid = pd.read_csv(source, dtype="string", keep_default_na=False)
    report = table_profile(geid)
    report.update({"source_sha256": digest(source), "damage_counts": counts(geid.DAMAGE_GRADE),
                   "district_counts": counts(geid.DISTRICT_ID), "source_references": counts(geid.REFERENCE),
                   "coordinates": "No latitude/longitude fields. District/VDC-municipality/ward IDs only.",
                   "distinct_building_id_strings": int(geid.BUILDING_ID.nunique()),
                   "duplicate_building_id_rows_after_first": int(geid.BUILDING_ID.duplicated().sum()),
                   "identifier_warning": "Source CSV IDs are rounded scientific-notation strings; cannot recover unique building keys by changing parser dtype."})
    district_categories = {c["value"]: c for c in definitions["dist"]["categories"]}
    comparisons = [{"district_code": k, "district_name": district_categories[k]["label"], "geid_rows": v,
                    "official_ddi_rows": district_categories[k]["frequency"],
                    "exact_aggregate_match": v == district_categories[k]["frequency"]}
                   for k, v in report["district_counts"].items()]
    pd.DataFrame(comparisons).to_csv(output / "nepal_district_reconciliation.csv", index=False)
    official_summary = {k: v for k, v in official.items() if k != "variables"}
    official_summary.update({"locally_acquired_building_microdata": False,
                            "damage_categories_from_metadata": definitions["dm_grade"]["categories"],
                            "coordinate_variables": [v["name"] for v in official["variables"] if any(s in (v["name"] + " " + v["label"]).lower() for s in ("latitude", "longitude", "gps"))],
                            "access_status": "Official metadata acquired. Get Microdata endpoint returned an empty body; terms specify public-use access from CBS premises. Obtain authorized Building file from NSO; no microdata download link was exposed.",
                            "license": "NSO research-use conditions; no redistribution, no reidentification or identifying linkage. Aggregate reporting only."})
    result = {"geid": report, "official_hrhrs": official_summary, "district_reconciliation": comparisons,
              "provenance_conclusion": "Strong evidence for an HRHRS-derived 11-district projection: every district count agrees exactly, shared survey variables and five damage grades, GEID cites the NPC/KLL earthquake portal. Official full-file row-level equivalence remains unverified; do not add GEID and HRHRS as independent observations.",
              "geid_license": "GEM CC BY-NC-SA 4.0; retain upstream-source restrictions and do not redistribute Nepal row data from this project."}
    write_json(output / "nepal.json", result)
    return result


def ingest_rc616(output):
    path = RAW / "rc616/454_866_Priority_I_data.csv"
    raw = pd.read_csv(path, dtype="string", keep_default_na=False)
    frame = raw.add_prefix("raw__")
    frame["building_id"] = "rc616_" + raw["Experiment ID"]
    frame["source_record"] = pd.Series(range(1, len(raw) + 1), dtype="Int64")
    frame["event_id"] = raw["Case Name"].map(EVENTS).astype("string")
    if frame.event_id.isna().any() or frame.building_id.duplicated().any():
        raise ValueError("Unknown RC616 event or duplicate experiment ID")
    frame["source_case_id"] = raw["Case ID"]
    for source, target in [("Latitude", "latitude"), ("Longitude", "longitude"), ("No. \nFloors", "floors")]:
        frame[target] = pd.to_numeric(raw[source].replace("", pd.NA), errors="raise").astype("Float64")
    frame["structural_damage_label"] = raw["Structural \nDamage"]
    frame["structural_damage_ordinal"] = frame.structural_damage_label.map(RC_DAMAGE).astype("Int8")
    frame.to_parquet(output / "rc616.parquet", index=False)
    report = table_profile(raw)
    report.update({"source_sha256": digest(path), "event_counts": counts(frame.event_id),
                   "structural_damage_counts": counts(frame.structural_damage_label),
                   "damage_ordinal_counts": counts(frame.structural_damage_ordinal),
                   "unmapped_damage_labels": counts(frame.loc[frame.structural_damage_ordinal.isna(), "structural_damage_label"]),
                   "mapping": RC_DAMAGE, "coordinates": coordinate_profile(frame),
                   "event_coordinate_counts": {event: coordinate_profile(group) for event, group in frame.groupby("event_id")},
                   "masonry_damage_counts": counts(raw["Masonry Wall\nDamage"]),
                   "advertised_rows": 616, "retained_rows": len(frame), "removed_rows": 0,
                   "count_discrepancy": "Published description says 616; archived CSV has 619 unique Experiment IDs/event-case pairs. Reason not established; all records retained.",
                   "license": "CC BY-SA 3.0; Sim et al. (2016), DOI 10.7277/ACX0-DG18"})
    metadata = pd.read_csv(RAW / "rc616/454_866_Priority_I_metadata.csv", keep_default_na=False)
    metadata.to_csv(output / "rc616_variables.csv", index=False)
    write_json(output / "rc616.json", report)
    return report


def haversine_distances(left, right):
    """All-pairs metres for small engineering surveys; input columns lat, lon."""
    a, b = np.radians(np.asarray(left, dtype=float)), np.radians(np.asarray(right, dtype=float))
    if not np.isfinite(a).all() or not np.isfinite(b).all():
        raise ValueError("Distance input must contain finite coordinates")
    delta = a[:, None, :] - b[None, :, :]
    h = np.sin(delta[:, :, 0] / 2) ** 2 + np.cos(a[:, None, 0]) * np.cos(b[None, :, 0]) * np.sin(delta[:, :, 1] / 2) ** 2
    return 6371008.8 * 2 * np.arcsin(np.sqrt(np.clip(h, 0, 1)))


def audit_aci(output):
    path = RAW / "turkiye_aci133/ACI133v3.xlsx"
    workbook = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    sheet, formula_sheet = workbook.active, formulas.active
    rows = list(sheet.values)
    if sheet.title != "Sheet1" or len(rows[1]) != 106:
        raise ValueError("ACI133 workbook layout changed")
    columns = [f"c{i:03d}" for i in range(1, 107)]
    raw = pd.DataFrame(rows[3:], columns=columns)
    dictionary = []
    for i, col in enumerate(columns, start=1):
        group = "building" if i <= 36 else "building_site" if i <= 49 else "kriging" if i <= 86 else "USGS"
        dictionary.append({"column": col, "excel_column": openpyxl.utils.get_column_letter(i),
                           "source_header": rows[1][i - 1], "source_units": rows[2][i - 1], "group": group})
    pd.DataFrame(dictionary).to_csv(output / "aci133_variables.csv", index=False)
    audit_cells, red_rows = [], set()
    for row in formula_sheet.iter_rows(min_row=4):
        for cell in row:
            red = bool(cell.font.color and cell.font.color.type == "rgb" and cell.font.color.rgb == "FFFF0000")
            if red:
                red_rows.add(cell.row)
            if cell.data_type == "f" or red:
                audit_cells.append({"cell": cell.coordinate, "source_formula": cell.value if cell.data_type == "f" else None,
                                    "red_font_author_supplementation": red, "cached_value": sheet[cell.coordinate].value})
    write_json(output / "aci133_source_cells.json", audit_cells)
    # Source columns can mix numbers with literal "NaN" tokens. Preserve those
    # tokens as strings rather than inventing numeric values or dropping rows.
    frame = raw.astype("string")
    frame["building_id"] = "aci133_" + raw.c001.astype("string")
    frame["source_excel_row"] = pd.Series(range(4, len(raw) + 4), dtype="Int64")
    frame["latitude"], frame["longitude"] = raw.c037, raw.c038
    frame["damage_label_5class"] = raw.c021
    frame["damage_ordinal_5class"] = raw.c021.map(ACI_DAMAGE).astype("Int8")
    frame["has_author_supplemented_cells"] = frame.source_excel_row.isin(red_rows)
    if frame.building_id.duplicated().any() or frame.damage_ordinal_5class.isna().any():
        raise ValueError("ACI133 IDs or damage vocabulary changed; inspect source")
    frame.to_parquet(output / "aci133.parquet", index=False)
    existing = pd.read_parquet(ROOT / "data/processed/turkiye_2023/buildings.parquet")
    distances = haversine_distances(frame[["latitude", "longitude"]], existing[["latitude", "longitude"]].astype(float))
    candidates = np.argwhere(distances <= 100)
    pd.DataFrame([{"aci133_id": frame.iloc[i].building_id, "survey_id": existing.iloc[j].building_id,
                   "distance_m": float(distances[i, j]), "candidate_only_not_verified_match": True}
                  for i, j in candidates]).to_csv(output / "turkiye_overlap_candidates_100m.csv", index=False)
    overlap = {str(t): {"pairs": int((distances <= t).sum()), "aci133_records": int((distances.min(axis=1) <= t).sum()),
                       "survey_records": int((distances.min(axis=0) <= t).sum())} for t in (0, 10, 25, 50, 100)}
    report = table_profile(raw)
    report.update({"source_sha256": digest(path), "coordinates": coordinate_profile(frame), "unique_source_ids": int(raw.c001.nunique()),
                   "literal_NaN_tokens_by_column": {c: int(raw[c].astype("string").eq("NaN").sum()) for c in raw if raw[c].astype("string").eq("NaN").any()},
                   "damage_5class_counts": counts(raw.c021), "damage_3class_counts": counts(raw.c022),
                   "damage_5class_mapping": ACI_DAMAGE, "author_supplemented_red_cells": sum(c["red_font_author_supplementation"] for c in audit_cells),
                   "rows_with_author_supplementation": len(red_rows), "formula_cells": sum(c["source_formula"] is not None for c in audit_cells),
                   "missing_formula_cache_cells": sum(c["source_formula"] is not None and c["cached_value"] is None for c in audit_cells),
                   "duplicate_source_headers": {k: v for k, v in Counter(rows[1]).items() if v > 1},
                   "header_warning": "Repeated M7.8 labels and inconsistent M7.5/M7.7 and M6.3/M6.4/M6.7/M6.8 labels are preserved by physical column and source block; event assignments are not silently corrected.",
                   "overlap_threshold_metres": overlap, "nearest_distance_min_m": float(distances.min()),
                   "overlap_conclusion": "Zero exact coordinate matches does not prove independence. Nearby records and shared-site coordinates require building-level adjudication. Do not claim 559+242 independent buildings.",
                   "license": "CC BY 4.0 (Zenodo API record 13386343); Hariri-Ardebili (2024). Author-completed cells are not fresh field observations."})
    write_json(output / "aci133.json", report)
    workbook.close()
    formulas.close()
    return report, frame, existing


def safe_extract_zip(archive, destination):
    destination = Path(destination).resolve()
    for member in archive.infolist():
        target = (destination / member.filename).resolve()
        if not target.is_relative_to(destination) or "\\" in member.filename or ((member.external_attr >> 16) & 0o170000) == 0o120000:
            raise ValueError(f"Unsafe archive member: {member.filename}")
    archive.extractall(destination)


def audit_spatial(output, engineering):
    import geopandas as gpd
    import pyogrio
    import rasterio
    from bs4 import BeautifulSoup

    archive_path = RAW / "turkiye_2023_context/2023Turkey_earthquake_data.zip"
    if digest(archive_path, "md5") != "f3a3a9982cf63bd954616a898c7416b9":
        raise ValueError("Spatial archive checksum mismatch")
    extract = archive_path.parent / "extracted"
    with zipfile.ZipFile(archive_path) as archive:
        members = [{"name": m.filename, "bytes": m.file_size, "compressed_bytes": m.compress_size, "crc32": m.CRC} for m in archive.infolist()]
        if not all((extract / m.filename).exists() and (m.is_dir() or (extract / m.filename).stat().st_size == m.file_size) for m in archive.infolist()):
            safe_extract_zip(archive, extract)
    write_json(output / "spatial_archive_inventory.json", members)
    points = {name: gpd.GeoDataFrame(geometry=gpd.points_from_xy(frame.longitude, frame.latitude), crs="EPSG:4326")
              for name, frame in engineering.items()}
    vectors = []
    for path in sorted([*extract.rglob("*.shp"), *extract.rglob("*.gmt")]):
        info = pyogrio.read_info(path)
        entry = {"path": str(path.relative_to(extract)), "features": info["features"], "fields": list(info["fields"]),
                 "dtypes": list(info["dtypes"]), "crs": info["crs"], "bounds": info["total_bounds"], "geometry_type": info["geometry_type"]}
        if info["features"] < 100_000:
            frame = pyogrio.read_dataframe(path)
            entry["features"] = len(frame)  # GMT drivers can return -1 (unknown) in metadata.
            entry["field_inventory"] = table_profile(frame.drop(columns="geometry"))["columns"]
            entry["category_counts"] = {c: counts(frame[c]) for c in frame if c != "geometry" and frame[c].nunique() <= 100}
            entry["invalid_geometries"] = int((~frame.geometry.is_valid).sum())
            entry["empty_geometries"] = int(frame.geometry.is_empty.sum())
            entry["duplicate_geometries_after_first"] = int(frame.geometry.to_wkb().duplicated().sum())
            entry["engineering_point_intersections"] = {name: int(gpd.sjoin(p, frame.to_crs(4326)[["geometry"]], how="inner", predicate="intersects").index.nunique()) for name, p in points.items()}
            entry["inspection_scope"] = "Full feature/attribute scan"
        else:
            entry["inspection_scope"] = "Full declared feature count/schema/bounds from shapefile metadata; no full 4.4M-geometry validity or attribute scan"
        vectors.append(entry)
    rasters = []
    for path in sorted(extract.rglob("*.tif")):
        with rasterio.open(path) as dataset:
            valid_count, minimum, maximum = 0, None, None
            histogram = np.zeros(256, dtype=np.int64) if dataset.dtypes[0] == "uint8" else None
            for _, window in dataset.block_windows(1):
                values = dataset.read(1, window=window, masked=True).compressed()
                values = values[np.isfinite(values)]
                if len(values):
                    valid_count += len(values)
                    minimum = float(values.min()) if minimum is None else min(minimum, float(values.min()))
                    maximum = float(values.max()) if maximum is None else max(maximum, float(values.max()))
                    if histogram is not None:
                        histogram += np.bincount(values, minlength=256)
            rasters.append({"file": path.name, "rows": dataset.height, "columns": dataset.width, "bands": dataset.count,
                            "dtype": dataset.dtypes, "crs": dataset.crs.to_string(), "bounds": list(dataset.bounds),
                            "pixel_size_degrees": dataset.res, "nodata": dataset.nodata, "units": dataset.units,
                            "scales": dataset.scales, "offsets": dataset.offsets, "tags": dataset.tags(),
                            "valid_pixels": valid_count, "masked_or_nonfinite_pixels": dataset.width * dataset.height - valid_count,
                            "min": minimum, "max": maximum,
                            "value_counts": {str(i): int(n) for i, n in enumerate(histogram) if n} if histogram is not None else None,
                            "inspection_scope": "Full band-1 pixel scan; values retained as encoded, units not inferred"})
    kml_layers = []
    namespace = {"k": "http://www.opengis.net/kml/2.2"}
    for path in sorted(extract.rglob("*.kmz")):
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("doc.kml"))
            placemarks = root.findall(".//k:Placemark", namespace)
            overlays = root.findall(".//k:GroundOverlay", namespace)
            attributes = []
            for placemark in placemarks:
                html = BeautifulSoup(placemark.findtext("k:description", default="", namespaces=namespace), "html.parser")
                record = {}
                for row in html.find_all("tr"):
                    cells = row.find_all("td", recursive=False)
                    if len(cells) == 2 and not any(c.find("table") for c in cells):
                        record[cells[0].get_text(strip=True)] = cells[1].get_text(strip=True)
                attributes.append(record)
            layer = {"file": path.name, "placemarks": len(placemarks), "ground_overlays": len(overlays),
                     "category_counts": {c: counts(pd.DataFrame(attributes)[c]) for c in pd.DataFrame(attributes) if pd.DataFrame(attributes)[c].nunique() <= 100},
                     "fields": list(pd.DataFrame(attributes).columns),
                     "overlay_extents": [{tag: float(overlay.findtext(f"k:LatLonBox/k:{tag}", namespaces=namespace)) for tag in ("north", "south", "east", "west")} for overlay in overlays]}
            layer["embedded_images"] = []
            for name in archive.namelist():
                if not name.endswith(".png"):
                    continue
                # Inventory IHDR only: huge rendered overlays need not be decoded.
                with archive.open(name) as stream:
                    header = stream.read(33)
                if header[:8] != b"\x89PNG\r\n\x1a\n" or header[12:16] != b"IHDR":
                    raise ValueError("Invalid embedded PNG header")
                width, height, depth, color_type, _, _, _ = struct.unpack(">IIBBBBB", header[16:29])
                layer["embedded_images"].append({"name": name, "width": width, "height": height,
                                                  "bit_depth": depth, "png_color_type": color_type})
            kml_layers.append(layer)
    report = {"archive_bytes": archive_path.stat().st_size, "archive_sha256": digest(archive_path),
              "uncompressed_bytes": sum(m["bytes"] for m in members), "members": len(members), "vectors": vectors,
              "rasters": rasters, "kmz": kml_layers, "license": "CC BY 4.0 per Zenodo API record 18437501. Retain original-source attributions; record license alone does not settle every upstream product's terms.",
              "label_warning": "Separate products, not additive independent building labels. Visual files lack an explicit damage-grade field (Join_Count is not a grade). ARIA files are rendered PNG overlays, not calibrated numeric DPM rasters. Jindires is in Syria. Raster physical units/codebook are absent."}
    write_json(output / "spatial.json", report)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    nepal = audit_nepal(args.output_dir)
    print("Nepal metadata and GEID audit complete", flush=True)
    rc = ingest_rc616(args.output_dir)
    aci, aci_frame, survey = audit_aci(args.output_dir)
    print("RC616 ingestion and ACI133 overlap audit complete", flush=True)
    spatial = audit_spatial(args.output_dir, {"survey_559": survey, "aci133_242": aci_frame})
    write_json(args.output_dir / "summary.json", {"nepal": nepal, "rc616": rc, "aci133": aci, "spatial": spatial})
    print(f"Audit complete: {args.output_dir}. No datasets pooled; no models fitted.")


if __name__ == "__main__":
    main()
