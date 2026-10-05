"""Offline, auditable ingestion: python -m src.data.turkiye [--input workbook]."""

import argparse
import hashlib
import json
from datetime import date, datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pyarrow

from .turkiye_schema import COLUMNS, DAMAGE_MAP, SCHEMA, dictionary

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = ROOT / "data/reference/turkiye_2023/Jaiswal_2023TurkiyeEQ_us6000jllz_field_str_damage_data.xlsx"
DEFAULT_OUTPUT = ROOT / "data/processed/turkiye_2023"
SHEET = "all_field_data"


def write_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")


def read_workbook(path):
    """Read literal values; report empty worksheet rows, reject formulas/schema drift."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != [SHEET]:
            raise ValueError(f"Expected only sheet {SHEET!r}; found {workbook.sheetnames}")
        sheet = workbook[SHEET]
        rows = sheet.iter_rows()
        header = [c.value if c.value is not None else f"Unnamed: {i}" for i, c in enumerate(next(rows))]
        expected = [c.source for c in COLUMNS]
        if header != expected:
            raise ValueError(f"Source columns changed; update dictionary explicitly. Found: {header}")
        records, source_rows, empty_rows = [], [], []
        for row_num, cells in enumerate(rows, start=2):
            if any(c.data_type in ("f", "e") for c in cells):
                raise ValueError(f"Formula or Excel error in row {row_num}; inspect source before ingesting")
            values = [c.value for c in cells]
            if all(v is None for v in values):
                empty_rows.append(row_num)
            else:
                records.append(values)
                source_rows.append(row_num)
        if not records:
            raise ValueError("Workbook contains no survey records")
        return pd.DataFrame(records, columns=header, dtype=object), source_rows, empty_rows
    finally:
        workbook.close()


def raw_text(value):
    if value is None or value is pd.NA:
        return pd.NA
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def normalize(series):
    """Trim exterior whitespace, treating only empty values as null, never 'None'."""
    return series.astype("string").str.strip().replace("", pd.NA)


def canonicalize(raw, source_rows):
    raw = raw.reset_index(drop=True)
    if list(raw.columns) != [c.source for c in COLUMNS]:
        raise ValueError("Raw columns must match the source dictionary exactly")
    if len(raw) != len(source_rows) or len(set(source_rows)) != len(source_rows):
        raise ValueError("Source rows must uniquely identify every input record")
    out = pd.DataFrame({
        "building_id": pd.Series([f"turkiye_2023_r{r:04d}" for r in source_rows], dtype="string"),
        "source_row": pd.Series(source_rows, dtype="Int64"),
    })
    issues = []

    def record(mask, field, kind):
        for i in out.index[mask.fillna(False)]:
            issues.append({"building_id": out.at[i, "building_id"], "source_row": int(out.at[i, "source_row"]),
                           "field": field, "kind": kind, "raw_value": raw_text(raw.at[i, field])})

    preserved = {}
    for col in COLUMNS:
        values = raw[col.source].map(raw_text).astype("string")
        preserved[f"raw__{col.source}"] = values
        clean = normalize(values)
        if col.dtype == "Float64":
            parsed = pd.to_numeric(clean, errors="coerce").astype("Float64")
            record(clean.notna() & parsed.isna(), col.source, "numeric_parse_error")
            if col.transform == "exp":
                with np.errstate(over="ignore", under="ignore", invalid="ignore"):
                    parsed = np.exp(parsed)
            out[col.canonical] = parsed
        elif col.dtype == "datetime64[ns]":
            out[col.canonical] = pd.to_datetime(clean, format="ISO8601", errors="coerce")
            record(clean.notna() & out[col.canonical].isna(), col.source, "date_parse_error")
        else:
            out[col.canonical] = clean
    out["damage_grade"] = out.damage_label.map(DAMAGE_MAP).astype("Int8")
    record(out.damage_label.notna() & out.damage_grade.isna(), "damage_condition", "unmapped_damage_label")
    out["has_damage_and_structure"] = (out.damage_grade.notna() & out.structure_type.notna()).astype("boolean")
    out = pd.concat([out, pd.DataFrame(preserved)], axis=1)
    validate_schema(out)
    return out, issues


def validate_schema(table):
    for col, dtype in SCHEMA.items():
        if col not in table or str(table[col].dtype) != dtype:
            raise ValueError(f"{col} must exist with dtype {dtype}")
    for col in COLUMNS:
        if f"raw__{col.source}" not in table or str(table[f"raw__{col.source}"].dtype) != "string":
            raise ValueError(f"Missing preserved source field: {col.source}")
    if table.building_id.isna().any() or table.building_id.duplicated().any():
        raise ValueError("building_id must be non-null and unique")
    if table.source_row.isna().any() or table.source_row.duplicated().any() or (table.source_row < 2).any():
        raise ValueError("source_row must be unique, non-null, and >= 2")
    expected_ids = table.source_row.map(lambda r: f"turkiye_2023_r{r:04d}").astype("string")
    if not table.building_id.equals(expected_ids):
        raise ValueError("building_id does not match source_row")
    expected_damage = table.damage_label.map(DAMAGE_MAP).astype("Int8")
    if not table.damage_grade.equals(expected_damage):
        raise ValueError("damage_grade must match explicit damage mapping")
    eligible = (table.damage_grade.notna() & table.structure_type.notna()).astype("boolean")
    if not table.has_damage_and_structure.equals(eligible):
        raise ValueError("has_damage_and_structure is inconsistent")


def valid_coordinates(table):
    return (table.latitude.between(-90, 90) & table.longitude.between(-180, 180)
            & ~((table.latitude == 0) & (table.longitude == 0))).fillna(False)


def counts(series):
    values = series.astype("string").fillna("<missing>").value_counts().sort_index()
    return {str(k): int(v) for k, v in values.items()}


def quality_report(table, parse_issues):
    """Report issues without changing the table. Missingness/duplicates are warnings."""
    invalid, impossible = [], []

    def add(target, mask, field, rule):
        for i in table.index[mask.fillna(False)]:
            target.append({"building_id": table.at[i, "building_id"], "source_row": int(table.at[i, "source_row"]),
                           "field": field, "value": str(table.at[i, field]), "rule": rule})

    for lat, lon in [("latitude", "longitude"), ("sample_latitude", "sample_longitude")]:
        for field, limit in [(lat, 90), (lon, 180)]:
            add(invalid, table[field].notna() & ~table[field].between(-limit, limit), field, f"finite and within [-{limit}, {limit}]")
        add(invalid, (table[lat] == 0) & (table[lon] == 0), lat, "(0, 0) is invalid for this Türkiye survey")
    for col in COLUMNS:
        if col.dtype != "Float64" or col.canonical in ("latitude", "longitude", "sample_latitude", "sample_longitude"):
            continue
        values = table[col.canonical]
        bad = values.notna() & ~np.isfinite(values)
        rule = "finite"
        if col.canonical in ("floors", "basement_floors"):
            minimum = 1 if col.canonical == "floors" else 0
            bad |= (values < minimum) | (values % 1 != 0)
            rule += f", integer >= {minimum}"
        elif col.canonical == "MMI":
            bad |= ~values.between(1, 12) & values.notna()
            rule += ", in [1, 12]"
        elif col.canonical in ("area", "vs30") or col.transform == "exp":
            bad |= values <= 0
            rule += ", positive"
        else:
            bad |= values < 0
            rule += ", nonnegative"
        add(impossible, bad, col.canonical, rule)
    coordinate_rows = table[table.latitude.notna() & table.longitude.notna()]
    duplicates = []
    for (lat, lon), group in coordinate_rows.groupby(["latitude", "longitude"], sort=True):
        if len(group) > 1:
            duplicates.append({"latitude": float(lat) if np.isfinite(lat) else str(lat),
                               "longitude": float(lon) if np.isfinite(lon) else str(lon), "count": len(group),
                               "building_ids": group.building_id.tolist(), "damage_labels": counts(group.damage_label)})
    raw_cols = [f"raw__{c.source}" for c in COLUMNS]
    exact = table.duplicated(raw_cols, keep=False)
    missing = {c: {"count": int(table[c].isna().sum()), "fraction": float(table[c].isna().mean())} for c in table}
    return {
        "row_count": len(table), "has_damage_and_structure_count": int(table.has_damage_and_structure.sum()),
        "missingness": missing, "damage_class_counts": counts(table.damage_grade),
        "damage_label_counts": counts(table.damage_label), "structure_type_counts": counts(table.structure_type),
        "city_counts": counts(table.city), "unique_coordinate_pairs": len(coordinate_rows.drop_duplicates(["latitude", "longitude"])),
        "duplicate_coordinate_groups": duplicates, "duplicate_coordinate_row_count": sum(d["count"] for d in duplicates),
        "exact_duplicate_row_ids": table.loc[exact, "building_id"].tolist(),
        "exact_duplicate_excess_count": int(table.duplicated(raw_cols).sum()),
        "missing_coordinate_row_ids": table.loc[table.latitude.isna() | table.longitude.isna(), "building_id"].tolist(),
        "invalid_coordinates": invalid, "impossible_values": impossible, "parse_and_mapping_issues": parse_issues,
        "error_count": len(invalid) + len(impossible) + len(parse_issues),
        "policy": "All survey records retained; no imputation, deduplication, winsorization, or spatial jitter. Empty worksheet rows explicitly inventoried in manifest.",
    }


def ingest(source=DEFAULT_INPUT, output=DEFAULT_OUTPUT):
    source, output = Path(source), Path(output)
    raw, rows, empty_rows = read_workbook(source)
    table, issues = canonicalize(raw, rows)
    report = quality_report(table, issues)
    output.mkdir(parents=True, exist_ok=True)
    table.to_parquet(output / "buildings.parquet", index=False, engine="pyarrow", compression="zstd")
    manifest = {
        "schema_version": 1, "source_filename": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "sheet": SHEET,
        "survey_record_count": len(table), "retained_record_count": len(table), "removed_record_count": 0,
        "empty_worksheet_rows": empty_rows, "damage_mapping": DAMAGE_MAP,
        "building_id_policy": "turkiye_2023_r + zero-padded Excel row. Identifies a survey record, not a verified unique physical building; stable for unchanged workbook row order.",
        "raw_policy": "Every source cell retained as raw__<header> string (dates ISO 8601, numbers Python scalar representation); raw whitespace preserved. Original workbook remains authoritative for Excel formatting/types.",
        "canonical_policy": "Trim outer text whitespace; only blank/whitespace missing. Numeric/date parse failures become null with original value and an issue. exp(log-intensity) to g or cm/s; uncertainty columns unchanged.",
        "parquet_sha256": hashlib.sha256((output / "buildings.parquet").read_bytes()).hexdigest(),
        "versions": {"pandas": pd.__version__, "numpy": np.__version__, "openpyxl": openpyxl.__version__, "pyarrow": pyarrow.__version__},
    }
    write_json(output / "manifest.json", manifest)
    write_json(output / "quality_report.json", report)
    write_json(output / "source_dictionary.json", dictionary())
    return table, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    table, report = ingest(args.input, args.output_dir)
    print(f"Retained {len(table)} records; {report['has_damage_and_structure_count']} with damage and structure. "
          f"Duplicate-coordinate rows: {report['duplicate_coordinate_row_count']}; errors: {report['error_count']}.")
    print(f"Parquet, dictionary, manifest and quality report: {args.output_dir}")
    if report["error_count"]:
        raise SystemExit("Validation errors found. Records retained; inspect quality_report.json before using them.")


if __name__ == "__main__":
    main()
